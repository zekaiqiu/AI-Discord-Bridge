"""read_file tool for the chat tool-loop (read_file_tool + haihub wiring).

The bug: the API runners' only tool was run_bash, which returns TEXT — so
PDFs (no pdftotext in the sandbox), Office docs, audio and video were
invisible to text-only models (glm / deepseek / minimax). read_file turns
any attachment into model-consumable text, running chat-backend-side where
pypdf and the Gemini multimodal key live.

Under test:
  * read_file_tool.read_file — text passthrough, PDF text-layer extraction,
    Office (docx/pptx/xlsx) zip+XML parsing, archive listing, directory
    listing, binary identification, and confinement to the allowed root.
  * haihub_runner._readfile_container_path / _readfile_allowed_root — a
    container path (/workspace/.attachments/<sid>/<name>) maps back to the
    chat-side attachments dir; escape attempts are refused.
  * read_file is registered in the tool schema.
"""
from __future__ import annotations

import base64
import tarfile
import zipfile
from pathlib import Path

import pytest

import haihub_runner
import read_file_tool


@pytest.fixture
def att(tmp_path: Path) -> Path:
    d = tmp_path / "attachments"
    d.mkdir()
    return d


# ---------------------------------------------------------------------------
# read_file_tool.read_file
# ---------------------------------------------------------------------------

def test_text_file_passthrough(att: Path):
    (att / "notes.txt").write_text("line one\nline two", encoding="utf-8")
    out = read_file_tool.read_file("notes.txt", allowed_root=att)
    assert "line one" in out and "line two" in out
    assert "notes.txt" in out


def test_missing_file(att: Path):
    out = read_file_tool.read_file("nope.txt", allowed_root=att)
    assert "no such file" in out


def test_rejects_empty_path(att: Path):
    assert "without a 'path'" in read_file_tool.read_file("", allowed_root=att)


def test_escape_outside_root_refused(att: Path):
    out = read_file_tool.read_file("/etc/hostname", allowed_root=att)
    assert "outside the readable area" in out or "no such file" in out


def test_dotdot_escape_refused(att: Path):
    secret = att.parent / "secret.txt"
    secret.write_text("top secret", encoding="utf-8")
    out = read_file_tool.read_file("../secret.txt", allowed_root=att)
    assert "top secret" not in out


def test_directory_listing(att: Path):
    (att / "a.txt").write_text("x")
    (att / "b.txt").write_text("yy")
    out = read_file_tool.read_file(".", allowed_root=att)
    assert "directory listing" in out and "a.txt" in out and "b.txt" in out


def test_office_docx_text(att: Path):
    pytest.importorskip("docx")
    from docx import Document
    doc = Document()
    doc.add_paragraph("The quarterly forecast improved markedly.")
    doc.save(str(att / "report.docx"))
    out = read_file_tool.read_file("report.docx", allowed_root=att)
    assert "quarterly forecast improved" in out


def test_office_xlsx_text(att: Path):
    pytest.importorskip("openpyxl")
    import openpyxl
    wb = openpyxl.Workbook()
    ws = wb.active
    ws["A1"] = "ticker"
    ws["B1"] = "price"
    ws["A2"] = "WZT"
    ws["B2"] = 123.45
    wb.save(str(att / "data.xlsx"))
    out = read_file_tool.read_file("data.xlsx", allowed_root=att)
    assert "WZT" in out and "123.45" in out


def test_zip_listing(att: Path):
    with zipfile.ZipFile(att / "bundle.zip", "w") as z:
        z.writestr("inner/readme.txt", "hello")
        z.writestr("inner/code.py", "print(1)")
    out = read_file_tool.read_file("bundle.zip", allowed_root=att)
    assert "zip archive" in out and "inner/code.py" in out


def test_tar_listing(att: Path):
    inner = att / "f.txt"
    inner.write_text("data")
    with tarfile.open(att / "bundle.tar", "w") as t:
        t.add(inner, arcname="f.txt")
    out = read_file_tool.read_file("bundle.tar", allowed_root=att)
    assert "tar archive" in out and "f.txt" in out


def test_pdf_text_layer(att: Path):
    pytest.importorskip("reportlab")
    from reportlab.pdfgen import canvas
    c = canvas.Canvas(str(att / "doc.pdf"))
    c.drawString(100, 750, "Revenue rose forty-two percent year over year.")
    c.save()
    out = read_file_tool.read_file("doc.pdf", allowed_root=att)
    assert "forty-two percent" in out


def test_binary_identified(att: Path):
    # Non-text bytes that aren't a known media type -> binary info path.
    (att / "blob.bin").write_bytes(bytes(range(256)) * 4)
    out = read_file_tool.read_file("blob.bin", allowed_root=att)
    assert "binary file" in out or "hex" in out


# ---------------------------------------------------------------------------
# haihub_runner path mapping + tool registration
# ---------------------------------------------------------------------------

def test_container_path_mapped_to_attachments():
    out = haihub_runner._readfile_container_path(
        "/workspace/.attachments/SID9/img.png", "SID9", "/data/attachments/SID9")
    assert out == "/data/attachments/SID9/img.png"


def test_bare_filename_mapped_to_attachments():
    out = haihub_runner._readfile_container_path("img.png", "SID9", "/data/attachments/SID9")
    assert out == "/data/attachments/SID9/img.png"


def test_unrelated_path_passthrough():
    out = haihub_runner._readfile_container_path("/workspace/other/x.txt", "SID9", "/data/attachments/SID9")
    assert out == "/workspace/other/x.txt"


def test_allowed_root_is_attachments_dir():
    assert haihub_runner._readfile_allowed_root("/data/attachments/SID9") == Path("/data/attachments/SID9")
    assert haihub_runner._readfile_allowed_root(None) is None


def test_read_file_registered_in_tool_schema():
    names = [t["function"]["name"] for t in haihub_runner._TOOLS]
    assert "read_file" in names and "run_bash" in names
